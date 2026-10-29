from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = '运营部工作内容与执行规范_RAGFlow版.docx'

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)

styles = doc.styles
styles['Normal'].font.name = 'PingFang SC'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'PingFang SC')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Normal'].paragraph_format.line_spacing = 1.15
for name, size, color in [('Title', 20, '1F4E79'), ('Heading 1', 15, '1F4E79'), ('Heading 2', 12.5, '2F75B5'), ('Heading 3', 11, '404040')]:
    st = styles[name]
    st.font.name = 'PingFang SC'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'PingFang SC')
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(9 if name != 'Title' else 0)
    st.paragraph_format.space_after = Pt(5)

def shade(p, fill='EAF3F8'):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)

def add_item(text, style='List Bullet'):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p

def add_rule(title, text, keywords=None):
    p = doc.add_paragraph()
    r = p.add_run(title + '：')
    r.bold = True
    p.add_run(text)
    if keywords:
        k = doc.add_paragraph()
        k.paragraph_format.left_indent = Cm(0.6)
        rr = k.add_run('检索关键词：')
        rr.bold = True
        k.add_run('、'.join(keywords))
        k.style = 'Normal'

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('运营部工作内容与执行规范')
# 清除 Word Title 样式可能带入的标题下边框，避免形成装饰性横线。
title_ppr = title._p.get_or_add_pPr()
title_bdr = OxmlElement('w:pBdr')
for edge in ('top', 'left', 'bottom', 'right', 'between', 'bar', 'insideH', 'insideV'):
    el = OxmlElement('w:' + edge)
    el.set(qn('w:val'), 'nil')
    title_bdr.append(el)
title_ppr.append(title_bdr)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run('适用于全店保險管家、运营巡检、会员服务、部门协同和对外沟通').italic = True

intro = doc.add_paragraph()
intro.add_run('文档目的：').bold = True
intro.add_run('明确运营部在硬件环境、软性服务、部门协调监管和对外沟通方面的工作内容、执行要求及异常处理方式。各条目按独立知识点编排，便于按主题、岗位、场景和关键词检索。')
shade(intro)

doc.add_heading('1 硬件环境', level=1)
add_rule('1.1 场内温度和湿度', '持续关注场馆温度、湿度和空调设备运行情况，发现异常时及时记录并联系相关维修或设备人员处理。', ['温度', '湿度', '空调设备监测'])
add_rule('1.2 场内环境卫生及绿植', '检查场馆卫生、绿植状态和保洁质量；涉及保洁或绿植问题时，与对应厂商做好对接。', ['环境卫生', '绿植', '保洁', '绿植厂商'])
add_rule('1.3 广播及场内音乐', '按照场馆广播规范播放内容，确保音乐类型与经营场景相匹配，音量保持适中。', ['广播规范', '场内音乐', '音乐类型', '音量'])
add_rule('1.4 器械和组合器械', '检查有氧器械、组合器械、杠铃、哑铃、杠铃台板等器械，以及私教区辅助健身用具的状态，确保设备安全、可用。', ['有氧器械', '组合器械', '杠铃', '哑铃', '私教区'])
add_rule('1.5 洗浴设备及浴室环境', '保障热水正常供应，检查洗手液、沐浴液、拖鞋等备品是否充足，并确认更衣柜、吹风机等设施可以正常使用。', ['热水', '洗手液', '沐浴液', '拖鞋', '更衣柜', '吹风机'])

doc.add_heading('1.6 多部门巡检', level=2)
add_item('运营部每日巡检：按日检查场馆环境和运营状态，发现问题及时沟通维修部门。')
add_item('维修部每周巡检：每周检查设备和场馆设施，并将问题反馈给运营部复检。')
add_item('私教部每日巡检：每日检查私教区域和相关设施，发现问题及时报备运营部。')
add_item('会籍部每日带领客户参观：带领客户参观时同步关注场馆问题，发现问题后及时报备运营部。')

doc.add_heading('1.7 多部门监督与复检', level=2)
add_item('运营部复检：对巡检问题进行复核，确认整改是否完成。')
add_item('店长复检：对重点问题和整改结果进行复核。')
add_item('总部相关部门复检：总部相关部门根据需要对场馆问题进行复检。')
add_item('总部审计部复检：审计部门对相关管理和执行情况进行复检。')

doc.add_heading('2 软性服务', level=1)
doc.add_heading('2.1 会员及客户接待', level=2)
add_item('正常会员入场：核验会员身份或入场信息后，引导会员正常入场。')
add_item('非会员体验或来访：根据现场安排进行会籍分配，并引导体验、来访人员完成接待流程。')
add_item('职能部门来访：做好来访登记、接待和现场引导。')
add_item('电话咨询接待：准确解答咨询内容，无法即时确认时记录并转交相关人员。')
add_item('寻人、求职、基础客诉：按场馆接待规范进行登记、分流和反馈。')
add_item('私教会员咨询：引导会员对接私教部门，确保咨询事项得到回应。')
add_item('办理新会员入会：按照入会流程核验资料、完成登记并引导会员办理入会。')

doc.add_heading('2.2 会员各项费用的收款', level=2)
add_item('会籍卡费用：按照会员合同和收费标准收取会籍卡相关费用。')
add_item('私教课费用：按照私教合同和收费标准收取私教课程费用。')
add_item('租柜费用：按照租柜规则收取租柜费用。')
add_item('定金费用：记录定金金额、用途和后续处理状态。')
add_item('水吧商品费用：按照商品价格完成水吧商品收款。')
add_item('转店及补卡费用：根据转店、补卡规则核对信息后收取费用。')

doc.add_heading('2.3 门店各项费用缴纳及能耗控制', level=2)
add_item('能耗费用管理：关注热水费、水费、电费和天然气等能源费用，通过各项管理措施降低能耗费用。')
add_item('门店经营费用管理：管理房租、维修费、日用品、低值易耗品、固定资产、保洁费、绿植费、电话费、饮用水和商品成本费用，做到成本控制并及时缴纳。')

doc.add_heading('2.4 全店备品及商品的盘点', level=2)
add_item('采购、入库、出库、调拨、报废和盘点等流程必须符合标准，妥善保管资产，确保资产目录清晰。')
add_item('盘点结果必须准确，确保门店商品账目清晰、账实相符。')

doc.add_heading('2.5 客诉处理', level=2)
add_rule('客诉定义', '客户对企业产品质量或服务质量不满意而提出的书面或口头异议、抗议、索赔和要求解决问题等行为。', ['客户投诉', '异议', '抗议', '索赔'])
doc.add_heading('2.5.1 客诉类型', level=3)
add_item('销售类客诉：销售时各项权利义务条款解释不清，或销售时承诺后期未履行。')
add_item('服务类客诉：服务标准未达标，所承诺的服务项目未履行，或存在服务态度问题。')
add_item('会员自觉类客诉：会员对相关问题存在不正确的个人理解。')
doc.add_heading('2.5.2 客诉处理流程', level=3)
add_item('按照“隔离—倾听—沟通—上报—提方案—反馈—跟踪—复盘”的流程处理客诉。')
doc.add_heading('2.5.3 客诉处理原则', level=3)
add_item('及时性：收到客诉后及时响应和处理。')
add_item('同理心：先梳理情绪，再梳理事件，理解客户感受和诉求。')
add_item('会员和公司双赢：在符合公司制度和合同约定的前提下，兼顾会员体验与公司利益。')
doc.add_heading('2.5.4 常见情景处理', level=3)
add_item('会员带16岁以下儿童独自进店游泳：按照未成年人入场和安全管理要求处理，确认监护责任后再决定是否允许入场。')
add_item('会员由于个人原因要求退卡：核对合同、退卡条件和实际履约情况后，按制度办理或上报审批。')
add_item('门店洗浴多次不出热水：立即记录时间和影响范围，联系维修人员排查，并向运营负责人反馈处理进度。')

doc.add_heading('2.6 突发事件处理', level=2)
add_rule('突发事件定义', '指突然发生、造成或者可能造成严重社会危害，需要采取应急处置措施予以应对的自然灾害、事故灾难、公共卫生事件和社会安全事件。', ['突发事件', '应急处置', '自然灾害', '事故灾难', '公共卫生事件', '社会安全事件'])
doc.add_heading('2.6.1 会员受伤', level=3)
add_item('第一时间送医，并上报保险公司、上报领导，保存现场影像资料。')
doc.add_heading('2.6.2 丢失贵重物品', level=3)
add_item('广播寻物，查看监控，协助寻找；报案时上报领导，现场拍照并保存视频，同时向保险公司报备。')
doc.add_heading('2.6.3 突发停水停电', level=3)
add_item('查看是否有人员受伤，安抚现场会员；广播通知场内会员，微信通知即将来访会员；联系维修人员和后勤排查原因并维修。')
doc.add_heading('2.6.4 物品损坏或玻璃破碎', level=3)
add_item('检查是否有人员受伤，现场拍照、保存视频，现场摆放警示牌，协商赔偿事宜，并尽快维修。')
doc.add_heading('2.6.5 漏水', level=3)
add_item('被漏水区域：转移物品、拍照取证，联系商场运营及工程或物业进行维修，现场摆放警示牌。')
add_item('漏到楼下：到现场拍照取证，联系工程，确认责任。')
doc.add_heading('2.6.6 政府机关检查或电视台采访', level=3)
add_item('不签收任何资料，不作任何承诺，表示新来的、不知情；不主动要联系方式，弄清对方来意并确认对方身份，第一时间上报领导。')

doc.add_heading('3 部门协调监管', level=1)
doc.add_heading('3.1 会籍部', level=2)
add_item('轮排表填写：由运营填写并监督轮排。')
add_item('会籍合同：正确填写合同相关内容并正确提交；在 MPL 范围内，超出范围时确认是否提前完成审批。')
add_item('代会员缴费：有明确会员授权代缴费和代签字；会员合同保存在前台处并及时交给会员。')
add_item('会籍提交的延期开卡、停卡、转卡、转店和通店等申请：按标准提交申请，并确保提前审批完毕。')
add_item('定金合同：确认是否在七日内补余。')

doc.add_heading('3.2 私教部', level=2)
doc.add_heading('3.2.1 签课要求', level=3)
add_item('会员在系统上有预约，当天上课、当天签课。')
add_item('教练携带 PT Dairy。')
add_item('前台打印小票，核实合同编号、所属课程、上课时间和上课节数等课程信息，确认无误；三体硬件自动签课监督。')
add_item('由教练、会员、前台三方同时在场，由本人手写签字。')
add_item('一节课时长不得低于45分钟。')
add_item('签课的前提是必须上了课，否则视为欺骗行为。')
add_item('满足以上所有要求，则运营可在签课小票上签字，此节课生效。')
doc.add_heading('3.2.2 授课数量', level=3)
add_item('同种类课程每天只允许上一节，通过钉钉审批最多可增加到2节。')
add_item('会员有多种类课程，但只有一种课程可进行加课申请。')
doc.add_heading('3.2.3 代会员缴费', level=3)
add_item('有明确会员授权代缴费和代签字，会员合同保存在前台处并及时交给会员。')
doc.add_heading('3.2.4 课程服务', level=3)
add_item('课程延期、延期开课、更换教练、更换会员、过期激活、提前开课、特殊情况停课和超 MPL 规定外销售，按标准提交申请，并确保提前审批完毕。')
doc.add_heading('3.2.5 课程取消', level=3)
add_item('教练应告知运营经理，运营经理在约课时间的3天内（含3天）告知总部取消。')
doc.add_heading('3.2.6 重大问题', level=3)
add_item('关于黑私教、私自承诺、拼单、垫付、刷单等严重问题，运营有监督举报的责任和义务。')

doc.add_heading('3.3 团操部', level=2)
add_item('确保团操上课设备设施顺畅使用。')
add_item('记录每节操课上课人数，确保课程有效开展。')
add_item('监督团操教师迟到早退等纪律问题。')
add_item('将会员关于上课排课意见和建议及时反馈给团操经理。')

doc.add_heading('4 对外沟通', level=1)
add_item('4.1 甲方商场或写字间：沟通店内缴纳费用、店内装修硬件损坏报修、停车票相关问题，参与商场内部营营活动和广告牌广告位等事项。')
add_item('4.2 工商部门：处理客诉或举报。')
add_item('4.3 税务部门：核税。')
add_item('4.4 城管、派出所：处理外圈宣传的公关沟通。')
add_item('4.5 消防大队：配合日常消防检查的接待和沟通。')
add_item('4.6 各种证照办理：办理体育经营备案证、食品经营许可证、卫生许可证和高危行业许可证。')

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('运营部工作内容与执行规范').font.size = Pt(8)

doc.core_properties.title = '运营部工作内容与执行规范'
doc.core_properties.subject = '全店保險管家工作内容及执行要求'
doc.core_properties.author = '运营部'
doc.save(OUT)
print(OUT)
